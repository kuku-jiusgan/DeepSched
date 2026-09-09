import unittest
from datetime import date, datetime, timedelta
from unittest.mock import patch

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.core.database import Base
from app.models import Instrument, Project, Task, TaskExecutionSegment, TimeSlot, User
from app.schemas.schemas import InstrumentOperatorUtilization, UtilizationStats
from app.services.instrument_utilization_report_service import (
    _attach_operator_details,
    build_instrument_utilization_report,
    export_instrument_utilization_report,
    utilization_window,
)


class InstrumentUtilizationReportServiceTest(unittest.TestCase):
    def test_defaults_to_current_month_through_today(self):
        start, end = utilization_window(None, None)
        today = datetime.now().date()

        self.assertEqual(today.replace(day=1), start.date())
        self.assertEqual(today, end.date())
        self.assertLessEqual(end, datetime.now())
        self.assertEqual(datetime.min.time(), start.time())

    def test_historical_end_date_includes_full_day(self):
        start, end = utilization_window(date(2026, 9, 1), date(2026, 9, 7))

        self.assertEqual(datetime(2026, 9, 1), start)
        self.assertEqual(datetime(2026, 9, 8), end)

    def test_rejects_invalid_date_range(self):
        with self.assertRaisesRegex(ValueError, "开始日期不能晚于结束日期"):
            utilization_window(date(2026, 9, 2), date(2026, 9, 1))

    @patch("app.services.instrument_utilization_report_service.calculate_instrument_utilization")
    def test_uses_shared_utilization_calculation(self, calculate):
        calculate.return_value = []

        result = build_instrument_utilization_report(
            object(), date(2026, 9, 1), date(2026, 9, 7), 100,
        )

        self.assertEqual([], result)
        calculate.assert_called_once_with(
            unittest.mock.ANY,
            datetime(2026, 9, 1),
            datetime(2026, 9, 8),
            100,
        )

    def test_exports_utilization_rows(self):
        output = export_instrument_utilization_report([UtilizationStats(
            instrument_id=1,
            instrument_code="INST-001",
            instrument_name="液相色谱仪",
            total_available_hours=168,
            scheduled_hours=24,
            actual_run_hours=16,
            expected_utilization_rate=14.3,
            actual_utilization_rate=9.5,
            utilization_rate=9.5,
            buffer_consumed_rate=0,
            operators=[InstrumentOperatorUtilization(
                operator_id=7,
                operator_name="张三",
                planned_hours=12,
                actual_run_hours=8,
                utilization_rate=4.8,
            )],
        )])

        self.assertGreater(len(output.getvalue()), 1000)
        from openpyxl import load_workbook

        workbook = load_workbook(output)
        self.assertIn("人员明细", workbook.sheetnames)
        self.assertEqual("张三", workbook["人员明细"]["C2"].value)

    def test_operator_detail_ignores_cancelled_open_slot_and_uses_segment_operator(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        db = sessionmaker(bind=engine)()
        self.addCleanup(db.close)
        assignee = User(username="assignee", display_name="王福芳", role="技术员")
        operator = User(username="operator", display_name="刘文静", role="技术员")
        instrument = Instrument(code="ZBYY-002-0001", name="液质联用仪")
        project = Project(code="XM-001", name="测试项目")
        db.add_all([assignee, operator, instrument, project])
        db.flush()
        task = Task(
            project_id=project.id, name="方法开发", task_type="FFKF_001",
            assignee_id=assignee.id, status="completed",
        )
        db.add(task)
        db.flush()
        dirty_anchor = TimeSlot(
            task_id=task.id, instrument_id=instrument.id,
            plan_start=datetime(2026, 9, 1, 10), plan_end=datetime(2026, 9, 1, 10),
            actual_start=datetime(2026, 8, 27, 10), actual_end=None,
            status="cancelled", lifecycle_status="superseded",
        )
        executed_slot = TimeSlot(
            task_id=task.id, instrument_id=instrument.id,
            plan_start=datetime(2026, 9, 2, 10), plan_end=datetime(2026, 9, 2, 12),
            status="cancelled", lifecycle_status="superseded",
        )
        db.add_all([dirty_anchor, executed_slot])
        db.flush()
        db.add(TaskExecutionSegment(
            task_id=task.id, slot_id=executed_slot.id, instrument_id=instrument.id,
            operator_id=operator.id, started_at=datetime(2026, 9, 2, 10),
            ended_at=datetime(2026, 9, 2, 12), end_reason="completed",
        ))
        db.commit()
        rows = [UtilizationStats(
            instrument_id=instrument.id, instrument_code=instrument.code,
            instrument_name=instrument.name, total_available_hours=192,
            scheduled_hours=0, actual_run_hours=2, expected_utilization_rate=0,
            actual_utilization_rate=1, utilization_rate=1, buffer_consumed_rate=0,
        )]

        [result] = _attach_operator_details(
            db, rows, datetime(2026, 9, 1), datetime(2026, 9, 9),
        )

        self.assertEqual(["刘文静"], [item.operator_name for item in result.operators])
        self.assertEqual(2.0, result.operators[0].actual_run_hours)


if __name__ == "__main__":
    unittest.main()
