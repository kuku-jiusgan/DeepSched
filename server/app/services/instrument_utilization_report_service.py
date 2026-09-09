from __future__ import annotations

from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models import Task, TaskExecutionSegment, TaskNightRun, TimeSlot
from app.schemas.schemas import InstrumentOperatorUtilization, UtilizationStats
from app.services.instrument_utilization_service import (
    _effective_work_ranges,
    _intersections,
    _load_faults_by_instrument,
    _subtract_ranges,
    _covered_hours as _covered_work_hours,
    calculate_instrument_utilization,
)


def build_instrument_utilization_report(
    db,
    start_date: date | None,
    end_date: date | None,
    percent_scale: float,
) -> list[UtilizationStats]:
    window_start, window_end = utilization_window(start_date, end_date)
    rows = calculate_instrument_utilization(db, window_start, window_end, percent_scale)
    if not rows:
        return rows
    return _attach_operator_details(db, rows, window_start, window_end)


def _attach_operator_details(db, rows, window_start, window_end):
    instrument_ids = [row.instrument_id for row in rows if row.instrument_id is not None]
    slots = db.query(TimeSlot).filter(
        TimeSlot.instrument_id.in_(instrument_ids),
        (
            (TimeSlot.plan_end > window_start) & (TimeSlot.plan_start < window_end)
        ) | (
            TimeSlot.actual_start.isnot(None)
            & (TimeSlot.actual_start < window_end)
            & (TimeSlot.actual_end.is_(None) | (TimeSlot.actual_end > window_start))
        ),
    ).all()
    task_ids = {slot.task_id for slot in slots}
    tasks = db.query(Task).filter(Task.id.in_(task_ids)).all() if task_ids else []
    task_map = {task.id: task for task in tasks}
    slot_ids = [slot.id for slot in slots]
    segments = db.query(TaskExecutionSegment).filter(
        TaskExecutionSegment.slot_id.in_(slot_ids),
        TaskExecutionSegment.started_at < window_end,
        (TaskExecutionSegment.ended_at.is_(None) | (TaskExecutionSegment.ended_at > window_start)),
    ).all() if slot_ids else []
    nights = db.query(TaskNightRun).filter(
        TaskNightRun.instrument_id.in_(instrument_ids),
        TaskNightRun.lifecycle_status == "active",
        TaskNightRun.started_at < window_end,
        TaskNightRun.ended_at > window_start,
    ).all()
    slot_by_id = {slot.id: slot for slot in slots}
    segment_task_ids = {segment.task_id for segment in segments}
    segment_slot_ids = {segment.slot_id for segment in segments}
    by_instrument: dict[int, dict[tuple[int | None, str], dict[str, list[tuple]]]] = {}

    def add_range(instrument_id, operator_id, operator_name, kind, start, end):
        if not start or not end:
            return
        start = max(start, window_start)
        end = min(end, window_end)
        if end <= start:
            return
        person = by_instrument.setdefault(instrument_id, {}).setdefault(
            (operator_id, operator_name), {"planned": [], "actual": [], "night": []},
        )
        person[kind].append((start, end))

    def assignee(task):
        user = task.assignee if task else None
        return (getattr(user, "id", None), getattr(user, "display_name", None) or "未分配")

    for slot in slots:
        task = task_map.get(slot.task_id)
        operator_id, operator_name = assignee(task)
        if _slot_counts_as_planned(slot):
            add_range(slot.instrument_id, operator_id, operator_name, "planned", slot.plan_start, slot.plan_end)
        if slot.task_id not in segment_task_ids and _slot_counts_as_actual(slot):
            add_range(slot.instrument_id, operator_id, operator_name, "actual", slot.actual_start, slot.actual_end or window_end)
    for segment in segments:
        slot = slot_by_id.get(segment.slot_id)
        if not slot or (segment.ended_at is None and not _slot_can_have_open_actual(slot)):
            continue
        operator_id = segment.operator_id
        operator_name = getattr(segment.operator, "display_name", None) or "未分配"
        add_range(slot.instrument_id, operator_id, operator_name, "actual", segment.started_at, segment.ended_at or window_end)
    for night in nights:
        operator_name = getattr(night.operator, "display_name", None) or "未分配"
        add_range(night.instrument_id, night.operator_id, operator_name, "night", night.started_at, night.ended_at)

    details = {}
    effective_ranges = {
        row.instrument_id: _effective_work_ranges(db, window_start, window_end, row.instrument_id)
        for row in rows
    }
    fault_ranges = _load_faults_by_instrument(db, window_start, window_end)
    for instrument_id, people in by_instrument.items():
        details[instrument_id] = []
        denominator = next((row.total_available_hours for row in rows if row.instrument_id == instrument_id), 0)
        allocated_planned = _allocate_planned_hours(
            people, effective_ranges.get(instrument_id, []),
        )
        for (operator_id, operator_name), ranges in people.items():
            planned = allocated_planned.get((operator_id, operator_name), 0.0)
            actual_ranges = [
                *_intersections(ranges["actual"], effective_ranges.get(instrument_id, [])),
                *ranges["night"],
            ]
            actual = _covered_work_hours(_subtract_ranges(
                actual_ranges, fault_ranges.get(instrument_id, []),
            ))
            details[instrument_id].append(InstrumentOperatorUtilization(
                operator_id=operator_id,
                operator_name=operator_name,
                planned_hours=round(planned, 1),
                actual_run_hours=round(actual, 1),
                utilization_rate=round(actual / denominator * 100 if denominator else 0, 1),
            ))
        details[instrument_id].sort(key=lambda item: item.operator_name)
    return [row.model_copy(update={"operators": details.get(row.instrument_id, [])}) for row in rows]


def _covered_hours(ranges):
    return _covered_work_hours(ranges)


def _allocate_planned_hours(people, boundaries):
    """Allocate each physical planned interval once across operators.

    The instrument total is a union of time intervals.  Personnel rows must use
    the same physical-time basis; otherwise overlapping operators (or night
    plans) make their rows add up to more than the instrument total.
    """
    allocated: dict[tuple[int | None, str], float] = {}
    occupied = []
    entries = [
        (operator, interval)
        for operator, ranges in people.items()
        for interval in _intersections(ranges["planned"], boundaries)
    ]
    entries.sort(key=lambda item: (item[1][0], item[1][1], item[0][1], item[0][0] or -1))
    for operator, (start, end) in entries:
        remaining = _subtract_ranges([(start, end)], occupied)
        allocated[operator] = allocated.get(operator, 0.0) + _covered_hours(remaining)
        occupied.extend(remaining)
        occupied = _merge_ranges(occupied)
    return allocated


def _merge_ranges(ranges):
    if not ranges:
        return []
    ordered = sorted(ranges)
    merged = [ordered[0]]
    for start, end in ordered[1:]:
        if start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    return merged


def _slot_counts_as_planned(slot: TimeSlot) -> bool:
    return bool(
        slot.lifecycle_status == "active"
        and slot.status != "cancelled"
        and slot.plan_start
        and slot.plan_end
        and slot.plan_end > slot.plan_start
    )


def _slot_counts_as_actual(slot: TimeSlot) -> bool:
    if not slot.actual_start:
        return False
    if slot.actual_end is not None:
        return slot.actual_end > slot.actual_start
    return _slot_can_have_open_actual(slot)


def _slot_can_have_open_actual(slot: TimeSlot) -> bool:
    return slot.lifecycle_status == "active" and slot.status == "running"


def utilization_window(
    start_date: date | None,
    end_date: date | None,
) -> tuple[datetime, datetime]:
    today = datetime.now().date()
    start = start_date or today.replace(day=1)
    end = end_date or today
    if end > today:
        raise ValueError("筛选结束日期不能晚于当前日期")
    if start > end:
        raise ValueError("开始日期不能晚于结束日期")
    window_start = datetime.combine(start, time.min)
    window_end = datetime.now() if end == today else datetime.combine(end + timedelta(days=1), time.min)
    return window_start, window_end


def export_instrument_utilization_report(rows: list[UtilizationStats]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "仪器利用率"
    headers = [
        "仪器编码", "仪器名称", "可用时长(h)", "计划占用(h)", "实际运行(h)",
        "预期利用率(%)", "实际利用率(%)",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="165C4A")
    for row in rows:
        sheet.append([
            row.instrument_code or f"ID-{row.instrument_id}",
            row.instrument_name,
            row.total_available_hours,
            row.scheduled_hours,
            row.actual_run_hours,
            row.expected_utilization_rate,
            row.actual_utilization_rate,
        ])
    for column, width in {"A": 20, "B": 24, "C": 16, "D": 16, "E": 16, "F": 18, "G": 18}.items():
        sheet.column_dimensions[column].width = width
    detail = workbook.create_sheet("人员明细")
    detail.append(["仪器编码", "仪器名称", "执行人", "计划占用(h)", "实际运行(h)", "利用率(%)"])
    for cell in detail[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="165C4A")
    for row in rows:
        for operator in row.operators:
            detail.append([
                row.instrument_code or f"ID-{row.instrument_id}", row.instrument_name,
                operator.operator_name, operator.planned_hours,
                operator.actual_run_hours, operator.utilization_rate,
            ])
    for column, width in {"A": 20, "B": 24, "C": 20, "D": 16, "E": 16, "F": 16}.items():
        detail.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
